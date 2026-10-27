#!/usr/bin/env python3
"""Build a two-column image-and-description PRD as a Word (.docx) file.

Usage:
    python3 build_docx_prd.py prd-spec.json output.docx [--template company.dotx]

Spec (image paths are relative to the spec file):
    {
      "title": "批量导入 产品需求文档",
      "meta": "版本 v1.0 ｜ 2026/10/22",             # optional
      "page": {"orientation": "portrait"},          # optional: portrait | landscape
      "font": {"latin": "Arial", "east_asian": "微软雅黑", "size_pt": 10.5},  # optional
      "label_color": "C00000",                      # optional; null = black labels
      "heading_shift": -1,                          # optional; "##" -> Heading 1 by default
      "blocks": [
        {"type": "markdown", "text": "## 1. 项目背景\\n- ..."},
        {"type": "image", "path": "arch.png", "width_cm": 15, "caption": "图 1 整体架构"},
        {"type": "page_break"},
        {"type": "design_table",
         "headers": ["功能", "原型图与功能说明"], "left_ratio": 0.25,
         "rows": [{"hierarchy": "批量导入 / 任务管理", "name": "导入任务列表",
                   "images": ["shots/01.png"],
                   "sections": [{"label": "功能与效果", "text": "..."}]}]}
      ]
    }

Markdown support: headings, one paragraph per line, "-"/"1." lists with indentation,
**bold**, pipe tables, and standalone image lines ![alt](path). List numbers are
kept as written (Word does not renumber them).
"""

from __future__ import annotations

import argparse
import io
import json
import math
import re
import sys
import zipfile
from pathlib import Path

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.image.image import Image as DocxImage
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Emu, Pt, RGBColor
except ImportError:  # pragma: no cover
    print("ERROR: python-docx is required (pip install python-docx)", file=sys.stderr)
    sys.exit(2)

HEADING_RE = re.compile(r"^(#{1,6})\s+(.*\S)\s*$")
IMAGE_RE = re.compile(r"^!\[([^\]]*)\]\(\s*<?([^)>]+?)>?\s*\)\s*$")
LIST_RE = re.compile(r"^(\s*)([-*+•]|\d+[.)、])\s+(.*)$")
BOLD_RE = re.compile(r"(\*\*[^*]+\*\*)")
TABLE_SEP_RE = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
BULLETS = ("•", "◦", "▪")
CJK_CM_PER_PT = 0.0353  # one CJK glyph at 1pt is ~0.0353cm wide


class BuildError(Exception):
    pass


# ---------------------------------------------------------------- helpers

def set_run_fonts(run_or_style, latin: str | None, east_asian: str | None) -> None:
    font = run_or_style.font
    if latin:
        font.name = latin
    element = run_or_style.element if hasattr(run_or_style, "element") else run_or_style._element
    rpr = element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    if latin:
        rfonts.set(qn("w:ascii"), latin)
        rfonts.set(qn("w:hAnsi"), latin)
    if east_asian:
        rfonts.set(qn("w:eastAsia"), east_asian)


def shade(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def repeat_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def fixed_layout(table) -> None:
    tblpr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)


def set_widths(table, widths_cm: list[float]) -> None:
    """Set grid, column, cell and table widths so Word and LibreOffice agree."""
    table.autofit = False
    tblpr = table._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:type"), "dxa")
    tblw.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    for col, width in zip(table.columns, widths_cm):
        col.width = Cm(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)


def text_width_units(text: str) -> float:
    return sum(1.0 if ord(ch) > 0x2E80 else 0.55 for ch in text)


def ensure_style(doc, name: str, base: str = "Normal"):
    styles = doc.styles
    try:
        return styles[name]
    except KeyError:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        try:
            style.base_style = styles[base]
        except KeyError:
            pass
        return style


def style_or_none(doc, name: str):
    try:
        return doc.styles[name]
    except KeyError:
        return None


# ---------------------------------------------------------------- container

class Container:
    """Wraps the document body or a table cell so the same renderer works in both."""

    def __init__(self, target, is_cell: bool):
        self.target = target
        self.is_cell = is_cell
        self._fresh = is_cell  # a new cell already holds one empty paragraph

    def paragraph(self, style=None):
        if self._fresh and self.target.paragraphs and not self.target.paragraphs[0].text:
            self._fresh = False
            p = self.target.paragraphs[0]
            if style is not None:
                p.style = style
            return p
        self._fresh = False
        return self.target.add_paragraph(style=style)

    def table(self, rows: int, cols: int):
        self._fresh = False
        return self.target.add_table(rows=rows, cols=cols)


# ---------------------------------------------------------------- builder

class Builder:
    def __init__(self, spec: dict, base_dir: Path, template: Path | None):
        self.spec = spec
        self.base = base_dir
        self.errors: list[str] = []
        self.doc = self._open(template)
        self.label_color = spec.get("label_color", "C00000")
        self.heading_shift = int(spec.get("heading_shift", -1))
        font = spec.get("font") or {}
        self.latin = font.get("latin", "Arial" if not template else None)
        self.east = font.get("east_asian", "微软雅黑" if not template else None)
        self.size_pt = float(font.get("size_pt", 10.5))
        self._setup_page(spec.get("page"), template is not None)
        self._setup_styles(template is not None and not spec.get("font"))

    # -- setup
    def _open(self, template: Path | None):
        if not template:
            return Document()
        data = template.read_bytes()
        if template.suffix.lower() == ".dotx":
            src = zipfile.ZipFile(io.BytesIO(data))
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
                for item in src.infolist():
                    content = src.read(item.filename)
                    if item.filename == "[Content_Types].xml":
                        content = content.replace(
                            b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
                            b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml")
                    dst.writestr(item, content)
            data = buf.getvalue()
        doc = Document(io.BytesIO(data))
        body = doc.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)
        return doc

    def _setup_page(self, page: dict | None, has_template: bool) -> None:
        section = self.doc.sections[0]
        if not has_template:
            section.page_width, section.page_height = Cm(21), Cm(29.7)
            for side in ("left_margin", "right_margin"):
                setattr(section, side, Cm(2))
            section.top_margin = section.bottom_margin = Cm(2.2)
        if page and page.get("orientation") == "landscape" and section.orientation != WD_ORIENT.LANDSCAPE:
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = section.page_height, section.page_width
        self.usable_emu = section.page_width - section.left_margin - section.right_margin
        self.usable_h_emu = section.page_height - section.top_margin - section.bottom_margin

    def _setup_styles(self, keep_template_fonts: bool) -> None:
        normal = self.doc.styles["Normal"]
        if not keep_template_fonts:
            set_run_fonts(normal, self.latin, self.east)
            normal.font.size = Pt(self.size_pt)
            for name in ("Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
                st = style_or_none(self.doc, name)
                if st is not None:
                    set_run_fonts(st, self.latin, self.east)
        pf = normal.paragraph_format
        pf.space_after = Pt(3)
        pf.line_spacing = 1.2
        hier = ensure_style(self.doc, "PRD Hierarchy")
        hier.font.bold = True
        hier.font.size = Pt(self.size_pt)
        hier.paragraph_format.space_after = Pt(0)
        feat = ensure_style(self.doc, "PRD Feature")
        feat.font.bold = True
        feat.font.size = Pt(self.size_pt + 1.5)
        feat.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        feat.paragraph_format.space_after = Pt(0)
        cap = ensure_style(self.doc, "PRD Caption")
        cap.font.size = Pt(self.size_pt - 1.5)
        cap.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # -- inline
    def add_inline(self, paragraph, text: str, bold: bool = False) -> None:
        text = text.replace("\\*", "*").replace("`", "")
        for part in BOLD_RE.split(text):
            if not part:
                continue
            is_bold = part.startswith("**") and part.endswith("**") and len(part) > 4
            run = paragraph.add_run(part[2:-2] if is_bold else part)
            if is_bold or bold:
                run.bold = True

    # -- images
    def resolve(self, path: str) -> Path:
        p = Path(path)
        return p if p.is_absolute() else (self.base / p)

    def image_size_cm(self, path: Path, max_w_cm: float, max_h_cm: float) -> tuple[float, float]:
        img = DocxImage.from_file(str(path))
        w, h = img.px_width, img.px_height
        width = max_w_cm
        height = width * h / w
        if height > max_h_cm:
            height = max_h_cm
            width = height * w / h
        return width, height

    def add_picture(self, container: Container, path_str: str, max_w_cm: float, max_h_cm: float,
                    align_center: bool = True) -> float:
        path = self.resolve(path_str)
        p = container.paragraph()
        if align_center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if not path.exists():
            self.errors.append(f"image not found: {path_str}")
            run = p.add_run(f"[缺少图片：{path_str}]")
            run.font.color.rgb = RGBColor(0xC0, 0, 0)
            return 0.6
        width, height = self.image_size_cm(path, max_w_cm, max_h_cm)
        p.add_run().add_picture(str(path), width=Cm(width))
        p.paragraph_format.space_after = Pt(4)
        return height + 0.3

    # -- markdown
    def render_markdown(self, container: Container, text: str, width_cm: float, in_cell: bool) -> float:
        """Render markdown; returns an estimated rendered height in cm."""
        lines = text.replace("\r\n", "\n").split("\n")
        height = 0.0
        indent_stack: list[int] = []
        i = 0
        line_cm = self.size_pt * 1.2 * 0.0353 + 0.1
        chars_per_line = max(8.0, width_cm / (self.size_pt * CJK_CM_PER_PT))
        while i < len(lines):
            raw = lines[i]
            stripped = raw.strip()
            if not stripped:
                indent_stack = []
                i += 1
                continue
            if stripped.startswith("|"):
                block = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    block.append(lines[i].strip())
                    i += 1
                height += self.render_table(container, block, width_cm)
                indent_stack = []
                continue
            m = HEADING_RE.match(stripped)
            if m:
                indent_stack = []
                level = len(m.group(1)) + self.heading_shift
                if in_cell:
                    p = container.paragraph()
                    self.add_inline(p, m.group(2), bold=True)
                else:
                    level = min(max(level, 1), 4)
                    style = style_or_none(self.doc, f"Heading {level}")
                    p = container.paragraph(style=style)
                    self.add_inline(p, m.group(2), bold=style is None)
                height += line_cm * 1.4
                i += 1
                continue
            m = IMAGE_RE.match(stripped)
            if m:
                indent_stack = []
                height += self.add_picture(container, m.group(2), width_cm, self.max_image_h_cm(in_cell))
                i += 1
                continue
            m = LIST_RE.match(raw.expandtabs(4))
            if m:
                indent = len(m.group(1))
                while indent_stack and indent < indent_stack[-1]:
                    indent_stack.pop()
                if not indent_stack or indent > indent_stack[-1]:
                    indent_stack.append(indent)
                level = len(indent_stack) - 1
                marker = m.group(2)
                if marker in "-*+•":
                    marker = BULLETS[min(level, len(BULLETS) - 1)]
                p = container.paragraph()
                pf = p.paragraph_format
                pf.left_indent = Cm(0.55 * (level + 1) + (0.15 if len(marker) > 1 else 0))
                pf.first_line_indent = Cm(-0.55 - (0.15 if len(marker) > 1 else 0))
                pf.space_after = Pt(1)
                p.add_run(marker + " ")
                self.add_inline(p, m.group(3))
                usable = max(4.0, chars_per_line - 1.6 * (level + 1))
                height += line_cm * max(1, math.ceil(text_width_units(m.group(3)) / usable))
                i += 1
                continue
            indent_stack = []
            p = container.paragraph()
            self.add_inline(p, stripped)
            height += line_cm * max(1, math.ceil(text_width_units(stripped) / chars_per_line)) + 0.05
            i += 1
        return height

    def render_table(self, container: Container, block: list[str], width_cm: float) -> float:
        rows = [r for r in block if not TABLE_SEP_RE.match(r)]
        cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
        if not cells:
            return 0.0
        cols = max(len(r) for r in cells)
        table = container.table(len(cells), cols)
        grid = style_or_none(self.doc, "Table Grid")
        if grid is not None:
            table.style = grid
        for r_idx, row in enumerate(cells):
            for c_idx in range(cols):
                cell = table.cell(r_idx, c_idx)
                cell.width = Cm(width_cm / cols)
                p = cell.paragraphs[0]
                self.add_inline(p, row[c_idx] if c_idx < len(row) else "", bold=(r_idx == 0))
                p.paragraph_format.space_after = Pt(0)
                if r_idx == 0:
                    shade(cell, "F2F2F2")
        set_widths(table, [width_cm / cols] * cols)
        if len(cells) > 1:
            repeat_header(table.rows[0])
        if not container.is_cell:
            self.doc.add_paragraph()
        return 0.75 * len(cells) + 0.3

    def max_image_h_cm(self, in_cell: bool) -> float:
        page_h = Emu(self.usable_h_emu).cm
        return min(page_h * 0.7, 17.0) if in_cell else page_h * 0.8

    # -- blocks
    def build(self) -> None:
        doc = self.doc
        body = Container(doc, is_cell=False)
        if self.spec.get("title"):
            style = style_or_none(doc, "Title")
            p = body.paragraph(style=style)
            self.add_inline(p, self.spec["title"], bold=style is None)
        if self.spec.get("meta"):
            p = body.paragraph(style=doc.styles["PRD Caption"])
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.add_run(self.spec["meta"])
        page_w_cm = Emu(self.usable_emu).cm
        for n, block in enumerate(self.spec.get("blocks", []), start=1):
            kind = block.get("type")
            if kind == "markdown":
                self.render_markdown(body, block.get("text", ""), page_w_cm, in_cell=False)
            elif kind == "image":
                width = min(float(block.get("width_cm", page_w_cm)), page_w_cm)
                self.add_picture(body, block["path"], width, self.max_image_h_cm(False))
                if block.get("caption"):
                    body.paragraph(style=doc.styles["PRD Caption"]).add_run(block["caption"])
            elif kind == "page_break":
                body.paragraph().add_run().add_break(WD_BREAK.PAGE)
            elif kind == "design_table":
                self.design_table(block, page_w_cm)
            else:
                raise BuildError(f"block {n}: unknown type {kind!r}")

    def design_table(self, block: dict, page_w_cm: float) -> None:
        doc = self.doc
        headers = block.get("headers") or ["功能", "原型图与功能说明"]
        ratio = float(block.get("left_ratio", 0.25))
        left_cm = page_w_cm * ratio
        right_cm = page_w_cm - left_cm
        inner_right = right_cm - 0.45
        inner_left = left_cm - 0.45
        rows = block.get("rows") or []
        if not rows:
            raise BuildError("design_table has no rows")
        # keep the branch heading and intro on the same page as the table
        for p in reversed(doc.paragraphs[-4:]):
            if p.text.strip():
                p.paragraph_format.keep_with_next = True
        table = doc.add_table(rows=1, cols=2)
        grid = style_or_none(doc, "Table Grid")
        if grid is not None:
            table.style = grid
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        fixed_layout(table)
        for idx, title in enumerate(headers[:2]):
            cell = table.rows[0].cells[idx]
            cell.width = Cm(left_cm if idx == 0 else right_cm)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(title).bold = True
            shade(cell, "F2F2F2")
        repeat_header(table.rows[0])
        for r_num, row in enumerate(rows, start=1):
            if not row.get("name"):
                raise BuildError(f"design_table row {r_num}: missing name")
            cells = table.add_row().cells
            cells[0].width, cells[1].width = Cm(left_cm), Cm(right_cm)
            right = Container(cells[1], is_cell=True)
            est = 0.25
            for img in row.get("images") or []:
                est += self.add_picture(right, img, inner_right, self.max_image_h_cm(True))
            if row.get("image_note"):
                p = right.paragraph(style=doc.styles["PRD Caption"])
                p.add_run(row["image_note"])
                est += 0.5
            for section in row.get("sections") or []:
                if section.get("label"):
                    p = right.paragraph()
                    run = p.add_run(section["label"])
                    run.bold = True
                    if self.label_color:
                        run.font.color.rgb = RGBColor.from_string(self.label_color)
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(1)
                    est += 0.65
                est += self.render_markdown(right, section.get("text", ""), inner_right, in_cell=True)
            if row.get("body"):
                est += self.render_markdown(right, row["body"], inner_right, in_cell=True)
            if not (row.get("sections") or row.get("body")):
                self.errors.append(f"row {row['name']}: no explanation")
            # left cell: hierarchy on top, feature name roughly at the vertical middle
            left = Container(cells[0], is_cell=True)
            hier_cm = 0.0
            if row.get("hierarchy"):
                p = left.paragraph(style=doc.styles["PRD Hierarchy"])
                p.add_run(row["hierarchy"])
                per_line = max(4.0, inner_left / (self.size_pt * CJK_CM_PER_PT))
                hier_cm = 0.55 * max(1, math.ceil(text_width_units(row["hierarchy"]) / per_line))
            name_p = left.paragraph(style=doc.styles["PRD Feature"])
            name_p.add_run(row["name"])
            gap_cm = max(0.3, est / 2 - hier_cm - 0.35)
            name_p.paragraph_format.space_before = Pt(min(gap_cm * 28.35, 1500))
        set_widths(table, [left_cm, right_cm])
        doc.add_paragraph()  # keeps the next table from merging with this one

    def save(self, out: Path) -> None:
        self.doc.save(str(out))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("spec", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--template", type=Path, help="Company .dotx/.docx whose styles and page setup to use")
    args = parser.parse_args()
    try:
        spec = json.loads(args.spec.read_text(encoding="utf-8"))
        builder = Builder(spec, args.spec.resolve().parent, args.template)
        builder.build()
        args.output.parent.mkdir(parents=True, exist_ok=True)
        builder.save(args.output)
    except (OSError, json.JSONDecodeError, BuildError, KeyError, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2
    print(f"wrote {args.output}")
    for err in builder.errors:
        print(f"WARNING: {err}", file=sys.stderr)
    return 1 if builder.errors else 0


if __name__ == "__main__":
    sys.exit(main())
