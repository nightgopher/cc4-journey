#!/usr/bin/env python3
"""Tests for build_docx_prd.py and inspect_docx_prd.py (python-docx required)."""

import io
import json
import struct
import subprocess
import sys
import tempfile
import unittest
import zipfile
import zlib
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

SCRIPTS = Path(__file__).parents[1] / "scripts"
BUILD = SCRIPTS / "build_docx_prd.py"
INSPECT = SCRIPTS / "inspect_docx_prd.py"


def png(path: Path, w: int, h: int) -> None:
    raw = b"".join(b"\x00" + b"\xcc\xdd\xee" * w for _ in range(h))
    def chunk(tag, data):
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
    path.write_bytes(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
                     + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def run(script, *args):
    return subprocess.run([sys.executable, str(script), *map(str, args)], capture_output=True, text=True)


class DocxScriptTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        png(self.dir / "list.png", 288, 180)
        png(self.dir / "mobile.png", 120, 400)

    def tearDown(self):
        self.tmp.cleanup()

    def spec(self, **overrides):
        spec = {
            "title": "测试 PRD",
            "blocks": [
                {"type": "markdown", "text": "## 1. 项目背景\n- 背景说明\n## 3. 产品设计\n### 3.1 任务管理\n- 管理导入任务"},
                {"type": "design_table", "rows": [
                    {"hierarchy": "批量导入 / 任务管理", "name": "导入任务列表", "images": ["list.png"],
                     "sections": [{"label": "界面结构", "text": "1. 标题：批量导入\n2. \"任务状态\"筛选框\n   - 枚举值：全部、草稿\n   - 默认值：全部"}]},
                    {"hierarchy": "批量导入 / 任务管理", "name": "失败通知",
                     "sections": [{"label": "功能与效果", "text": "任务结束且存在失败条目时通知创建者。"}]},
                ]},
                {"type": "markdown", "text": "### 3.2 结果管理"},
                {"type": "design_table", "rows": [
                    {"hierarchy": "批量导入 / 结果管理", "name": "移动端结果页", "images": ["mobile.png"],
                     "body": "查看单条导入结果和失败原因。"}]},
            ],
        }
        spec.update(overrides)
        path = self.dir / "spec.json"
        path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
        return path

    def build(self, spec_path, *extra):
        out = self.dir / "out.docx"
        result = run(BUILD, spec_path, out, *extra)
        return result, out

    def test_build_and_inspect_pass(self):
        result, out = self.build(self.spec())
        self.assertEqual(result.returncode, 0, result.stderr)
        check = run(INSPECT, out, "--visual-row", "导入任务列表", "--expected-section", "任务管理",
                    "--expected-section", "结果管理", "--expected-tables", "2")
        self.assertEqual(check.returncode, 0, check.stdout)
        data = json.loads(check.stdout)
        self.assertEqual(data["row_count"], 3)
        rows = data["tables"][0]["rows"]
        self.assertEqual((rows[0]["hierarchy"], rows[0]["name"], rows[0]["images"]), ("批量导入 / 任务管理", "导入任务列表", 1))
        self.assertEqual(rows[1]["images"], 0)

    def test_column_widths_follow_ratio(self):
        _, out = self.build(self.spec())
        doc = Document(str(out))
        cols = [int(c.get(qn("w:w"))) for c in doc.tables[0]._tbl.tblGrid.findall(qn("w:gridCol"))]
        self.assertAlmostEqual(cols[0] / sum(cols), 0.25, places=2)

    def test_numbering_and_headings_kept_as_written(self):
        _, out = self.build(self.spec())
        doc = Document(str(out))
        heading = [p for p in doc.paragraphs if p.text == "1. 项目背景"][0]
        self.assertEqual(heading.style.name, "Heading 1")
        cell_text = [p.text for p in doc.tables[0].rows[1].cells[1].paragraphs]
        self.assertIn("1. 标题：批量导入", cell_text)
        self.assertTrue(any(t.startswith("◦ 枚举值") for t in cell_text), cell_text)

    def test_tall_image_is_height_capped(self):
        _, out = self.build(self.spec())
        doc = Document(str(out))
        shape = doc.inline_shapes[-1]
        self.assertLessEqual(shape.height.cm, 17.01)

    def test_missing_image_is_reported(self):
        spec = self.spec()
        data = json.loads(spec.read_text(encoding="utf-8"))
        data["blocks"][1]["rows"][0]["images"] = ["nope.png"]
        spec.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        result, out = self.build(spec)
        self.assertEqual(result.returncode, 1)
        self.assertIn("image not found", result.stderr)
        check = json.loads(run(INSPECT, out, "--visual-row", "导入任务列表").stdout)
        self.assertFalse(check["ok"])

    def test_landscape_option(self):
        _, out = self.build(self.spec(page={"orientation": "landscape"}))
        section = Document(str(out)).sections[0]
        self.assertGreater(section.page_width, section.page_height)

    def test_dotx_template_is_accepted(self):
        base = Document()
        base.add_paragraph("模板正文应被清除")
        buf = io.BytesIO()
        base.save(buf)
        src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
        dotx = self.dir / "company.dotx"
        with zipfile.ZipFile(dotx, "w") as dst:
            for item in src.infolist():
                content = src.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    content = content.replace(b"document.main+xml", b"template.main+xml")
                dst.writestr(item, content)
        result, out = self.build(self.spec(), "--template", dotx)
        self.assertEqual(result.returncode, 0, result.stderr)
        texts = [p.text for p in Document(str(out)).paragraphs]
        self.assertNotIn("模板正文应被清除", texts)

    def test_inspect_flags_text_before_image(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "功能", "原型图与功能说明"
        table.cell(1, 0).text = "导入任务列表"
        right = table.cell(1, 1)
        right.paragraphs[0].add_run("说明写在了图片前面，这是错误的顺序。")
        right.add_paragraph().add_run().add_picture(str(self.dir / "list.png"))
        path = self.dir / "bad.docx"
        doc.save(str(path))
        data = json.loads(run(INSPECT, path).stdout)
        self.assertFalse(data["ok"])
        self.assertTrue(any("before the image" in e for e in data["errors"]))

    def test_unknown_block_type_fails(self):
        result, _ = self.build(self.spec(blocks=[{"type": "chart"}]))
        self.assertEqual(result.returncode, 2)


if __name__ == "__main__":
    unittest.main()
